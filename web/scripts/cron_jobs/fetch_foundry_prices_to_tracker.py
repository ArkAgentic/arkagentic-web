#!/usr/bin/env python3
"""
Daily step 2: fetch Azure pricing for NEWLY appended models only.

Rules:
- Default scope is rows auto-appended by sync_llm_models_tracker.py (notes marker).
- If price_input/price_output already filled (including 0), skip fetch and do not email.
- Auto-matched rows: clear notes.
- Unmatched rows: notes=manual_price_input_required and optional email notification.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import smtplib
import subprocess
from email.message import EmailMessage
from pathlib import Path
from typing import Dict, Optional, Tuple

import requests
from openpyxl import load_workbook

PriceTriple = Tuple[float, Optional[float], float]  # input, cached_input, output per 1M

PROVIDER_SLUGS = [
    "aoai",
    "kimi",
    "deepseek",
    "cohere",
    "grok",
    "mistral-ai",
    "microsoft",
    "fireworks",
    "llama",
    "black-forest-labs",
]

SLUG_HINTS = {
    "kimi": ["kimi"],
    "deepseek": ["deepseek"],
    "cohere": ["cohere"],
    "grok": ["grok"],
    "mistral-ai": ["mistral"],
    "llama": ["llama"],
    "black-forest-labs": ["flux", "black forest"],
    "microsoft": ["mai-", "mai "],
    "aoai": ["gpt-", "text-embedding", "o1", "o3", "o4", "claude"],
}

AUTO_APPEND_NOTE_PREFIX = "auto-appended from Foundry deployment list"


def provider_slug_for_model(upstream_model_name: str) -> str:
    s = upstream_model_name.lower()
    for slug, hints in SLUG_HINTS.items():
        if any(h in s for h in hints):
            return slug
    return "aoai"


def pricing_url(slug: str) -> str:
    return f"https://azure.microsoft.com/en-us/pricing/details/ai-foundry-models/{slug}/"


def fetch_html(url: str) -> str:
    r = requests.get(url, timeout=30)
    r.raise_for_status()
    return r.text


def normalize_label(x: str) -> str:
    s = re.sub(r"\s+", " ", x.strip())
    return s.lower()


def extract_price_rows_from_html(html: str) -> Dict[str, PriceTriple]:
    out: Dict[str, PriceTriple] = {}
    row_pat = re.compile(r"<tr>\s*<td>([^<]+)</td>(.*?)</tr>", re.S | re.I)
    amount_pat = re.compile(r"data-amount='(\{.*?\})'", re.S | re.I)

    for m in row_pat.finditer(html):
        label = re.sub(r"\s+", " ", m.group(1)).strip()
        block = m.group(2)
        values = []
        for aj in amount_pat.findall(block):
            try:
                payload = json.loads(aj)
                regional = payload.get("regional", {}) if isinstance(payload, dict) else {}
                v = regional.get("us-east")
                if v is None and isinstance(regional, dict) and regional:
                    v = next(iter(regional.values()))
                if isinstance(v, (int, float)):
                    values.append(float(v))
            except Exception:
                continue

        if len(values) >= 2:
            inp = values[0]
            if len(values) == 2:
                cached = None
                outp = values[1]
            else:
                cached = values[1]
                outp = values[2]
            out[normalize_label(label)] = (inp, cached, outp)

    return out


def _canon(s: str) -> str:
    s = normalize_label(s)
    s = re.sub(r"\b(global|datazone)\b", "", s)
    s = re.sub(r"\s+", " ", s).strip()
    return re.sub(r"[^a-z0-9]+", "", s)


def best_match_price(upstream_model_name: str, rows: Dict[str, PriceTriple]) -> Optional[Tuple[str, PriceTriple]]:
    target = _canon(upstream_model_name)
    if not target:
        return None

    for k, v in rows.items():
        kc = _canon(k)
        if target == kc:
            return k, v
        if target and (target in kc or kc in target):
            return k, v

    tset = set([t for t in re.split(r"[^a-z0-9]+", normalize_label(upstream_model_name)) if t and t not in {"global", "datazone"}])
    best = None
    for k, v in rows.items():
        kset = set([t for t in re.split(r"[^a-z0-9]+", normalize_label(k)) if t and t not in {"global", "datazone"}])
        if not tset or not kset:
            continue
        inter = len(tset & kset)
        if inter == 0:
            continue
        score = inter / max(len(tset), 1)
        if best is None or score > best[0]:
            best = (score, k, v)
    if best and best[0] >= 0.6:
        return best[1], best[2]

    return None


def col_ref(col: int, row: int) -> str:
    letters = ""
    n = col
    while n:
        n, rem = divmod(n - 1, 26)
        letters = chr(65 + rem) + letters
    return f"{letters}{row}"


def headers_map(ws):
    return {str(ws.cell(1, c).value).strip(): c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}


def send_email(to_email: str, subject: str, body: str) -> Tuple[bool, str]:
    # Preferred: Azure Communication Services Email (prevents localhost sender spam issues)
    acs_conn = os.getenv("ACS_CONNECTION_STRING", "").strip()
    acs_sender = os.getenv("ACS_EMAIL_SENDER", "").strip()
    if acs_conn and acs_sender:
        try:
            cmd = [
                "az",
                "communication",
                "email",
                "send",
                "--connection-string",
                acs_conn,
                "--sender",
                acs_sender,
                "--subject",
                subject,
                "--to",
                to_email,
                "--text",
                body,
                "--wait-until",
                "completed",
                "-o",
                "json",
            ]
            p = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
            if p.returncode == 0:
                return True, "acs_email"
            return False, f"acs_email_failed: {(p.stderr or p.stdout).strip()[:300]}"
        except Exception as e:
            return False, f"acs_email_exception: {e}"

    # Secondary fallback: authenticated SMTP only (no localhost sender fallback)
    smtp_host = os.getenv("SMTP_HOST")
    smtp_port = int(os.getenv("SMTP_PORT", "587"))
    smtp_user = os.getenv("SMTP_USER")
    smtp_pass = os.getenv("SMTP_PASS")
    smtp_from = os.getenv("SMTP_FROM", "").strip()

    if smtp_host and smtp_user and smtp_pass and smtp_from:
        try:
            msg = EmailMessage()
            msg["Subject"] = subject
            msg["From"] = smtp_from
            msg["To"] = to_email
            msg.set_content(body)
            with smtplib.SMTP(smtp_host, smtp_port, timeout=30) as server:
                server.starttls()
                server.login(smtp_user, smtp_pass)
                server.send_message(msg)
            return True, "smtp"
        except Exception as e:
            return False, f"smtp_failed: {e}"

    return False, "no_email_transport_configured (set ACS_CONNECTION_STRING+ACS_EMAIL_SENDER)"

def has_manual_price(ws, headers, row: int) -> bool:
    p_in = ws.cell(row, headers["price_input_usd_per_1m"]).value
    p_out = ws.cell(row, headers["price_output_usd_per_1m"]).value

    def filled(v):
        if v is None:
            return False
        if isinstance(v, str):
            return v.strip() != ""
        return True

    return filled(p_in) and filled(p_out)


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--excel", default="LLM Models Tracker.xlsx")
    p.add_argument("--sheet", default="pricing_worksheet")
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--apply", action="store_true")
    p.add_argument("--all-candidates", action="store_true", help="Fetch for all rows missing prices, not only newly appended rows")
    p.add_argument("--notify-email", default="charles.zhang@arkagentic.com")
    p.add_argument("--no-email", action="store_true")
    args = p.parse_args()

    if not args.dry_run and not args.apply:
        raise SystemExit("Specify either --dry-run or --apply")

    excel_path = Path(args.excel).expanduser().resolve()
    if not excel_path.exists():
        raise SystemExit(f"Excel file not found: {excel_path}")

    wb = load_workbook(excel_path)
    ws = wb[args.sheet] if args.sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    headers = headers_map(ws)

    required = [
        "upstream_model_name",
        "ark_model_name",
        "cost_input_usd_per_1m",
        "cost_output_usd_per_1m",
        "price_input_usd_per_1m",
        "price_output_usd_per_1m",
        "profit_input_usd_per_1m",
        "profit_output_usd_per_1m",
        "gross_margin_pct",
        "notes",
    ]
    missing = [k for k in required if k not in headers]
    if missing:
        raise RuntimeError(f"Missing required columns: {missing}")

    page_rows: Dict[str, Dict[str, PriceTriple]] = {}
    page_errors: Dict[str, str] = {}
    for slug in PROVIDER_SLUGS:
        try:
            html = fetch_html(pricing_url(slug))
            page_rows[slug] = extract_price_rows_from_html(html)
        except Exception as e:
            page_rows[slug] = {}
            page_errors[slug] = str(e)

    updated = []
    unmatched = []
    skipped_price_present = []
    candidate_rows = []

    for r in range(2, ws.max_row + 1):
        upstream = ws.cell(r, headers["upstream_model_name"]).value
        ark = ws.cell(r, headers["ark_model_name"]).value
        note = ws.cell(r, headers["notes"]).value
        if not isinstance(upstream, str) or not upstream.strip():
            continue

        is_new_row = isinstance(note, str) and note.startswith(AUTO_APPEND_NOTE_PREFIX)
        if args.all_candidates:
            eligible = not has_manual_price(ws, headers, r)
        else:
            eligible = is_new_row

        if not eligible:
            continue

        candidate_rows.append(r)

        if has_manual_price(ws, headers, r):
            skipped_price_present.append({"row": r, "upstream_model_name": upstream, "ark_model_name": ark})
            continue

        slug = provider_slug_for_model(upstream)
        candidates = page_rows.get(slug, {})
        m = best_match_price(upstream, candidates)
        if not m and slug != "aoai":
            m = best_match_price(upstream, page_rows.get("aoai", {}))

        if not m:
            ws.cell(r, headers["notes"], "manual_price_input_required")
            unmatched.append({
                "row": r,
                "upstream_model_name": upstream,
                "ark_model_name": ark,
                "provider_slug": slug,
                "reason": "not_found_or_not_confident",
            })
            continue

        matched_label, (inp, cached, outp) = m
        if inp < 0 or outp < 0:
            ws.cell(r, headers["notes"], "manual_price_input_required")
            unmatched.append({
                "row": r,
                "upstream_model_name": upstream,
                "ark_model_name": ark,
                "provider_slug": slug,
                "reason": "invalid_negative_price",
            })
            continue

        ws.cell(r, headers["cost_input_usd_per_1m"], float(inp))
        ws.cell(r, headers["cost_output_usd_per_1m"], float(outp))

        # pricing/profit/gross-margin formulas
        ws.cell(r, headers["price_input_usd_per_1m"], f"={col_ref(headers['cost_input_usd_per_1m'], r)}*2")
        ws.cell(r, headers["price_output_usd_per_1m"], f"={col_ref(headers['cost_output_usd_per_1m'], r)}*2")
        ws.cell(r, headers["profit_input_usd_per_1m"], f"=IF({col_ref(headers['price_input_usd_per_1m'], r)}=\"\",\"\",{col_ref(headers['price_input_usd_per_1m'], r)}-{col_ref(headers['cost_input_usd_per_1m'], r)})")
        ws.cell(r, headers["profit_output_usd_per_1m"], f"=IF({col_ref(headers['price_output_usd_per_1m'], r)}=\"\",\"\",{col_ref(headers['price_output_usd_per_1m'], r)}-{col_ref(headers['cost_output_usd_per_1m'], r)})")
        ws.cell(r, headers["gross_margin_pct"], f"=IF(({col_ref(headers['price_input_usd_per_1m'], r)}+{col_ref(headers['price_output_usd_per_1m'], r)})=0,\"\",({col_ref(headers['price_input_usd_per_1m'], r)}+{col_ref(headers['price_output_usd_per_1m'], r)}-{col_ref(headers['cost_input_usd_per_1m'], r)}-{col_ref(headers['cost_output_usd_per_1m'], r)})/({col_ref(headers['price_input_usd_per_1m'], r)}+{col_ref(headers['price_output_usd_per_1m'], r)}))")

        ws.cell(r, headers["notes"], "")

        updated.append({
            "row": r,
            "upstream_model_name": upstream,
            "ark_model_name": ark,
            "provider_slug": slug,
            "matched_label": matched_label,
            "cost_input_usd_per_1m": inp,
            "cached_input_usd_per_1m": cached,
            "cost_output_usd_per_1m": outp,
        })

    email_attempted = False
    email_sent = False
    email_result = None

    # notify only if we actually attempted new rows and still have unmatched
    if args.apply and unmatched and candidate_rows and not args.no_email:
        email_attempted = True
        subject = f"[ArkAgentic] Manual price input required for {len(unmatched)} new model(s)"
        lines = [
            "New Foundry models were added but some prices could not be auto-filled.",
            "Please review and manually input costs in LLM Models Tracker.xlsx.",
            "",
            f"Excel: {excel_path}",
            "",
            "Rows requiring manual input:",
        ]
        for item in unmatched:
            lines.append(
                f"- row {item['row']}: upstream={item['upstream_model_name']} | ark={item['ark_model_name']} | slug={item['provider_slug']} | reason={item['reason']}"
            )
        email_sent, email_result = send_email(args.notify_email, subject, "\n".join(lines))

    if args.apply:
        wb.save(excel_path)

    print(json.dumps({
        "excel": str(excel_path),
        "dry_run": args.dry_run,
        "scope": "all_candidates" if args.all_candidates else "new_rows_only",
        "candidate_rows_count": len(candidate_rows),
        "provider_pages_loaded": {k: len(v) for k, v in page_rows.items()},
        "provider_page_errors": page_errors,
        "updated_count": len(updated),
        "updated_rows": updated,
        "unmatched_count": len(unmatched),
        "unmatched_rows": unmatched,
        "skipped_price_present_count": len(skipped_price_present),
        "skipped_price_present_rows": skipped_price_present,
        "email_attempted": email_attempted,
        "email_sent": email_sent,
        "email_result": email_result,
        "saved": args.apply,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
