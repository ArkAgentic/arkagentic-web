#!/usr/bin/env python3
"""
Daily step 3: sync ArkAgentic pricing from LLM Models Tracker.xlsx.

Targets:
1) DB model_pricing table (runtime source of truth)
2) DB routes/channels mapping sync via scripts/sync-azure-deployments.mjs

Rules:
- If either input/output price is missing, skip that model.
- 0 is valid (not missing).
- No static pricing-schema.ts rewrite here.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
from pathlib import Path
from typing import Dict, Optional, Tuple, List

from openpyxl import load_workbook


def headers_map(ws) -> Dict[str, int]:
    return {str(ws.cell(1, c).value).strip(): c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}


def col_ref_to_index(col_letters: str) -> int:
    n = 0
    for ch in col_letters:
        n = n * 26 + (ord(ch.upper()) - ord("A") + 1)
    return n


def parse_cell_numeric(ws, row: int, col: int, depth: int = 0) -> Optional[float]:
    if depth > 3:
        return None
    v = ws.cell(row, col).value

    if isinstance(v, (int, float)):
        return float(v)

    if isinstance(v, str):
        s = v.strip()
        if s == "":
            return None

        try:
            return float(s)
        except Exception:
            pass

        if s.startswith("="):
            t = s[1:].replace("$", "").replace(" ", "")

            m = re.fullmatch(r"([A-Z]+)(\d+)\*([0-9]*\.?[0-9]+)", t, re.I)
            if m:
                c_ref, r_ref, k = m.group(1), int(m.group(2)), float(m.group(3))
                base = parse_cell_numeric(ws, r_ref, col_ref_to_index(c_ref), depth + 1)
                return None if base is None else base * k

            m = re.fullmatch(r"([0-9]*\.?[0-9]+)\*([A-Z]+)(\d+)", t, re.I)
            if m:
                k, c_ref, r_ref = float(m.group(1)), m.group(2), int(m.group(3))
                base = parse_cell_numeric(ws, r_ref, col_ref_to_index(c_ref), depth + 1)
                return None if base is None else k * base

    return None


def load_tracker_rows(ws, headers) -> Tuple[List[dict], List[dict]]:
    updates = []
    skipped_missing_price = []

    for r in range(2, ws.max_row + 1):
        model_id = ws.cell(r, headers["ark_model_name"]).value
        upstream = ws.cell(r, headers["upstream_model_name"]).value
        if not isinstance(model_id, str) or not model_id.strip():
            continue
        model_id = model_id.strip()
        upstream_name = str(upstream).strip() if upstream is not None else ""

        in_price = parse_cell_numeric(ws, r, headers["price_input_usd_per_1m"])
        out_price = parse_cell_numeric(ws, r, headers["price_output_usd_per_1m"])

        if in_price is None or out_price is None:
            skipped_missing_price.append({"row": r, "ark_model_name": model_id})
            continue

        updates.append(
            {
                "row": r,
                "ark_model_name": model_id,
                "upstream_model_name": upstream_name,
                "price_input_usd_per_1m": in_price,
                "price_output_usd_per_1m": out_price,
                "input_price_per_1k": in_price / 1000.0,
                "output_price_per_1k": out_price / 1000.0,
            }
        )

    return updates, skipped_missing_price


def run_db_upsert(js_helper: Path, payload_path: Path, dry_run: bool, env: Dict[str, str]) -> dict:
    p = subprocess.run(
        ["node", str(js_helper), str(payload_path), "--dry-run" if dry_run else "--apply"],
        capture_output=True,
        text=True,
        env=env,
    )
    if p.returncode != 0:
        raise RuntimeError(f"model_pricing upsert failed (exit={p.returncode})\nSTDOUT:\n{p.stdout}\nSTDERR:\n{p.stderr}")
    return json.loads(p.stdout)


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--excel", default="LLM Models Tracker.xlsx")
    p.add_argument("--sheet", default="pricing_worksheet")
    p.add_argument("--db-sync-script", default="scripts/sync-azure-deployments.mjs")
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--apply", action="store_true")
    p.add_argument("--sync-db", action="store_true", default=True)
    p.add_argument("--no-sync-db", action="store_true")
    args = p.parse_args()

    if not args.dry_run and not args.apply:
        raise SystemExit("Specify either --dry-run or --apply")

    do_sync_db = args.sync_db and (not args.no_sync_db)

    excel_path = Path(args.excel).expanduser().resolve()
    if not excel_path.exists():
        raise SystemExit(f"Excel file not found: {excel_path}")

    db_script_path = Path(args.db_sync_script).expanduser().resolve()
    if do_sync_db and not db_script_path.exists():
        raise SystemExit(f"DB sync script not found: {db_script_path}")

    wb = load_workbook(excel_path)
    ws = wb[args.sheet] if args.sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    headers = headers_map(ws)

    required = ["ark_model_name", "upstream_model_name", "price_input_usd_per_1m", "price_output_usd_per_1m"]
    missing = [k for k in required if k not in headers]
    if missing:
        raise RuntimeError(f"Missing required columns: {missing}")

    updates, skipped_missing_price = load_tracker_rows(ws, headers)

    helper_js = Path(__file__).resolve().parent / "sync_model_pricing_db.mjs"
    payload_json = Path(__file__).resolve().parent / "_pricing_updates_payload.json"
    payload_json.write_text(json.dumps({"updates": updates}, ensure_ascii=False), encoding="utf-8")

    db_pricing_result = {
        "attempted": True,
        "ok": False,
        "details": None,
        "error": None,
    }

    try:
        db_pricing_result["details"] = run_db_upsert(helper_js, payload_json, dry_run=args.dry_run, env=dict(os.environ))
        db_pricing_result["ok"] = bool(db_pricing_result["details"].get("ok", False))
    except Exception as e:
        db_pricing_result["error"] = str(e)

    db_routes_result = {
        "enabled": do_sync_db,
        "attempted": False,
        "ok": False,
        "details": None,
        "error": None,
    }

    if do_sync_db:
        env = dict(os.environ)
        db_args = ["--json"]
        if args.dry_run:
            db_args.append("--dry-run")

        db_routes_result["attempted"] = True
        try:
            p2 = subprocess.run(["node", str(db_script_path)] + db_args, capture_output=True, text=True, env=env)
            if p2.returncode != 0:
                stderr_text = (p2.stderr or "").strip()
                stdout_text = (p2.stdout or "").strip()
                if stdout_text:
                    try:
                        parsed = json.loads(stdout_text)
                        db_routes_result["details"] = parsed
                        err = parsed.get("error") if isinstance(parsed, dict) else None
                        db_routes_result["error"] = f"exit={p2.returncode}; error={str(err)[:500]}"
                    except Exception:
                        db_routes_result["error"] = f"exit={p2.returncode}; stdout={stdout_text[:500]}"
                else:
                    db_routes_result["error"] = f"exit={p2.returncode}; stderr={stderr_text[:500]}"
            else:
                try:
                    db_routes_result["details"] = json.loads(p2.stdout)
                    db_routes_result["ok"] = bool(db_routes_result["details"].get("ok", True))
                except Exception:
                    db_routes_result["error"] = f"non-json output: {p2.stdout[:500]}"
        except Exception as e:
            db_routes_result["error"] = str(e)

    print(
        json.dumps(
            {
                "excel": str(excel_path),
                "db_sync_script": str(db_script_path),
                "dry_run": args.dry_run,
                "candidate_price_rows": len(updates),
                "skipped_missing_price_count": len(skipped_missing_price),
                "skipped_missing_price": skipped_missing_price,
                "db_pricing_sync": db_pricing_result,
                "db_routes_sync": db_routes_result,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
