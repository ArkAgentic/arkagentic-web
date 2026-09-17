#!/usr/bin/env python3
"""
Daily step 1: sync Azure Foundry deployments into LLM Models Tracker.xlsx.
- append new deployments
- prune models removed from Foundry
"""

from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List, Tuple

import requests
from openpyxl import load_workbook

REQUIRED_HEADERS = [
    "provider_channel_type",
    "upstream_model_name",
    "ark_model_name",
    "cost_input_usd_per_1m",
    "cost_output_usd_per_1m",
    "price_input_usd_per_1m",
    "price_output_usd_per_1m",
    "profit_input_usd_per_1m",
    "profit_output_usd_per_1m",
    "gross_margin_pct",
    "notes",
]


def run_cmd(cmd: List[str], timeout_s: int = 60) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, capture_output=True, text=True, timeout=timeout_s)


def ensure_az_logged_in(auto_login: bool = True) -> dict:
    p = run_cmd(["az", "account", "show", "-o", "json"], timeout_s=30)
    if p.returncode != 0 and auto_login:
        login = run_cmd(["az", "login", "--use-device-code", "-o", "json"], timeout_s=300)
        if login.returncode != 0:
            msg = (login.stderr or login.stdout or "").strip()
            raise RuntimeError(f"Azure CLI login failed: {msg}")
        p = run_cmd(["az", "account", "show", "-o", "json"], timeout_s=30)

    if p.returncode != 0:
        msg = (p.stderr or p.stdout or "").strip()
        raise RuntimeError(f"Azure CLI is not logged in: {msg}")
    return json.loads(p.stdout)


def run_az_deployments_via_arm(resource_group: str, account: str, subscription_id: str, bearer_token: str, timeout_s: int = 90) -> List[dict]:
    api_versions = [
        os.getenv("AZURE_ARM_DEPLOYMENTS_API_VERSION", "").strip(),
        "2024-10-01-preview",
        "2024-10-01",
        "2023-10-01-preview",
        "2023-05-01",
    ]
    api_versions = [v for v in api_versions if v]

    last_error = ""
    for api_version in api_versions:
        url = (
            f"https://management.azure.com/subscriptions/{subscription_id}"
            f"/resourceGroups/{resource_group}"
            f"/providers/Microsoft.CognitiveServices/accounts/{account}"
            f"/deployments?api-version={api_version}"
        )
        try:
            resp = requests.get(
                url,
                headers={"Authorization": f"Bearer {bearer_token}", "Content-Type": "application/json"},
                timeout=timeout_s,
            )
            text = resp.text
            data = json.loads(text)
            if resp.ok:
                items = data.get("value") if isinstance(data, dict) else None
                if isinstance(items, list):
                    return items
                if isinstance(data, list):
                    return data
                return []
            last_error = f"ARM {resp.status_code} [{api_version}]: {text[:300]}"
        except Exception as e:
            last_error = f"ARM exception [{api_version}]: {e}"

    raise RuntimeError(f"ARM deployment list failed: {last_error}")


def get_arm_token_from_imds(timeout_s: int = 10) -> str:
    client_id = (os.getenv("AZURE_CLIENT_ID") or "").strip()

    # Azure Container Apps / App Service style managed identity endpoint
    identity_endpoint = (os.getenv("IDENTITY_ENDPOINT") or "").strip()
    identity_header = (os.getenv("IDENTITY_HEADER") or "").strip()
    if identity_endpoint and identity_header:
        params = {
            "api-version": "2019-08-01",
            "resource": "https://management.azure.com/",
        }
        if client_id:
            params["client_id"] = client_id
        try:
            resp = requests.get(
                identity_endpoint,
                params=params,
                headers={"X-IDENTITY-HEADER": identity_header, "Metadata": "true"},
                timeout=timeout_s,
            )
            if resp.ok:
                data = resp.json()
                token = str(data.get("access_token") or "").strip()
                if token:
                    return token
        except Exception:
            pass

    # IMDS style endpoint
    params = {
        "api-version": "2018-02-01",
        "resource": "https://management.azure.com/",
    }
    if client_id:
        params["client_id"] = client_id

    try:
        resp = requests.get(
            "http://169.254.169.254/metadata/identity/oauth2/token",
            params=params,
            headers={"Metadata": "true"},
            timeout=timeout_s,
        )
        if not resp.ok:
            return ""
        data = resp.json()
        return str(data.get("access_token") or "").strip()
    except Exception:
        return ""


def get_arm_bearer_token(timeout_s: int = 30) -> Tuple[str, str]:
    env_token = (os.getenv("AZURE_ARM_BEARER_TOKEN") or "").strip()
    if env_token:
        return env_token, "env"

    imds_token = get_arm_token_from_imds(timeout_s=min(timeout_s, 10))
    if imds_token:
        return imds_token, "imds"

    token_proc = run_cmd([
        "az",
        "account",
        "get-access-token",
        "--resource",
        "https://management.azure.com/",
        "--query",
        "accessToken",
        "-o",
        "tsv",
    ], timeout_s=timeout_s)
    if token_proc.returncode == 0:
        token = token_proc.stdout.strip()
        if token:
            return token, "az"

    return "", "none"


def get_subscription_id(timeout_s: int = 30) -> str:
    env_sub = (os.getenv("AZURE_SUBSCRIPTION_ID") or os.getenv("AZURE_SUBSCRIPTION") or "").strip()
    if env_sub:
        return env_sub

    # derive from resource id if available in managed runtime
    az_sub_from_id = (os.getenv("AZURE_SUBSCRIPTION_RESOURCE_ID") or "").strip()
    if az_sub_from_id:
        m = re.search(r"/subscriptions/([^/]+)", az_sub_from_id, re.I)
        if m:
            return m.group(1).strip()

    # fallback: parse from UAMI resource id if provided
    uami_id = (os.getenv("AZURE_UAMI_RESOURCE_ID") or "").strip()
    if uami_id:
        m = re.search(r"/subscriptions/([^/]+)", uami_id, re.I)
        if m:
            return m.group(1).strip()

    sub_proc = run_cmd(["az", "account", "show", "--query", "id", "-o", "tsv"], timeout_s=timeout_s)
    if sub_proc.returncode == 0:
        return sub_proc.stdout.strip()

    return ""


def run_az_deployments(resource_group: str, account: str, timeout_s: int = 90) -> List[dict]:
    # first try CLI command
    cmd = [
        "az", "cognitiveservices", "account", "deployment", "list",
        "-g", resource_group, "-n", account, "-o", "json",
    ]
    p = run_cmd(cmd, timeout_s=timeout_s)
    if p.returncode == 0:
        payload = json.loads(p.stdout)
        if not isinstance(payload, list):
            raise RuntimeError("Unexpected Azure deployment response")
        return payload

    # fallback to ARM using env token / managed identity (IMDS) / az token
    token, token_source = get_arm_bearer_token(timeout_s=30)
    subscription_id = get_subscription_id(timeout_s=30)
    if not token or not subscription_id:
        raise RuntimeError(
            "Azure deployment list failed and ARM fallback prerequisites missing: "
            f"token_source={token_source}, has_subscription={bool(subscription_id)}; "
            f"cli_error={p.stderr.strip() or p.stdout.strip()}"
        )

    return run_az_deployments_via_arm(resource_group, account, subscription_id, token, timeout_s=timeout_s)


def slugify_model(name: str) -> str:
    slug = name.strip().lower()
    slug = slug.replace(".0-", "-").replace(".0", "")
    slug = re.sub(r"[^a-z0-9]+", "-", slug)
    slug = re.sub(r"-+", "-", slug).strip("-")
    return slug


def infer_ark_model_id(upstream_model_name: str) -> str:
    special = {
        "Cohere-embed-v3-multilingual": "ark-cohere-embed-v3",
        "Cohere-rerank-v4.0-fast": "ark-cohere-rerank-v4-fast",
        "Cohere-rerank-v4.0-pro": "ark-cohere-rerank-v4-pro",
    }
    if upstream_model_name in special:
        return special[upstream_model_name]
    return f"ark-{slugify_model(upstream_model_name)}"


def col_ref(col: int, row: int) -> str:
    letters = ""
    n = col
    while n:
        n, rem = divmod(n - 1, 26)
        letters = chr(65 + rem) + letters
    return f"{letters}{row}"


def header_index_map(ws) -> Dict[str, int]:
    header_to_col: Dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        v = ws.cell(1, col).value
        if isinstance(v, str) and v.strip():
            header_to_col[v.strip()] = col
    missing = [h for h in REQUIRED_HEADERS if h not in header_to_col]
    if missing:
        raise RuntimeError(f"Missing required headers in sheet '{ws.title}': {missing}")
    return header_to_col


def existing_upstream_models(ws, col_upstream: int) -> Dict[str, int]:
    out: Dict[str, int] = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, col_upstream).value
        if isinstance(v, str) and v.strip():
            out[v.strip()] = r
    return out


def append_model_rows(ws, h: Dict[str, int], additions: List[Tuple[str, str, str]]) -> int:
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    start_row = ws.max_row + 1
    for i, (provider, upstream, ark) in enumerate(additions):
        r = start_row + i
        ws.cell(r, h["provider_channel_type"], provider)
        ws.cell(r, h["upstream_model_name"], upstream)
        ws.cell(r, h["ark_model_name"], ark)
        ws.cell(r, h["cost_input_usd_per_1m"], None)
        ws.cell(r, h["cost_output_usd_per_1m"], None)
        ws.cell(r, h["price_input_usd_per_1m"], None)
        ws.cell(r, h["price_output_usd_per_1m"], None)

        c_cost_in = h["cost_input_usd_per_1m"]
        c_cost_out = h["cost_output_usd_per_1m"]
        c_price_in = h["price_input_usd_per_1m"]
        c_price_out = h["price_output_usd_per_1m"]
        c_profit_in = h["profit_input_usd_per_1m"]
        c_profit_out = h["profit_output_usd_per_1m"]
        c_margin = h["gross_margin_pct"]

        ws.cell(r, c_profit_in, f'=IF({col_ref(c_price_in, r)}="","",{col_ref(c_price_in, r)}-{col_ref(c_cost_in, r)})')
        ws.cell(r, c_profit_out, f'=IF({col_ref(c_price_out, r)}="","",{col_ref(c_price_out, r)}-{col_ref(c_cost_out, r)})')
        ws.cell(
            r,
            c_margin,
            f'=IF(({col_ref(c_price_in, r)}+{col_ref(c_price_out, r)})=0,"",({col_ref(c_price_in, r)}+{col_ref(c_price_out, r)}-{col_ref(c_cost_in, r)}-{col_ref(c_cost_out, r)})/({col_ref(c_price_in, r)}+{col_ref(c_price_out, r)}))',
        )
        ws.cell(r, h["notes"], f"auto-appended from Foundry deployment list @ {now}")
    return len(additions)


def prune_rows(ws, row_indexes: List[int]) -> int:
    for r in sorted(row_indexes, reverse=True):
        ws.delete_rows(r, 1)
    return len(row_indexes)


def main() -> None:
    parser = argparse.ArgumentParser(description="Sync Foundry deployments into LLM Models Tracker.xlsx")
    parser.add_argument("--excel", default="LLM Models Tracker.xlsx")
    parser.add_argument("--sheet", default="pricing_worksheet")
    parser.add_argument("--resource-group", default="AI-Resources-RG")
    parser.add_argument("--account", default="ArkAgentic")
    parser.add_argument("--provider", default="azure_openai")
    parser.add_argument("--include-nonrunning", action="store_true")
    parser.add_argument("--prune-missing", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--no-auto-login", action="store_true")
    args = parser.parse_args()

    account_info = {"id": None, "name": None}
    if not args.no_auto_login:
        account_info = ensure_az_logged_in(auto_login=True)
    else:
        # in non-interactive runtimes (e.g. Azure jobs), allow fallback via ARM token if available
        p = run_cmd(["az", "account", "show", "-o", "json"], timeout_s=30)
        if p.returncode == 0:
            try:
                account_info = json.loads(p.stdout)
            except Exception:
                account_info = {"id": None, "name": None}

    excel_path = Path(args.excel).expanduser().resolve()
    if not excel_path.exists():
        raise SystemExit(f"Tracker file not found: {excel_path}")

    deployments = run_az_deployments(args.resource_group, args.account)

    normalized: Dict[str, dict] = {}
    for d in deployments:
        name = str(d.get("name", "")).strip()
        if not name:
            continue
        state = str((d.get("properties") or {}).get("deploymentState") or "")
        if (not args.include_nonrunning) and state and state.lower() != "running":
            continue
        normalized[name] = d

    wb = load_workbook(excel_path)
    ws = wb[args.sheet] if args.sheet in wb.sheetnames else wb[wb.sheetnames[0]]

    h = header_index_map(ws)
    existing_map = existing_upstream_models(ws, h["upstream_model_name"])
    existing_set = set(existing_map.keys())
    foundry_set = set(normalized.keys())

    additions: List[Tuple[str, str, str]] = []
    for upstream_model in sorted(foundry_set, key=lambda s: s.lower()):
        if upstream_model in existing_set:
            continue
        additions.append((args.provider, upstream_model, infer_ark_model_id(upstream_model)))

    removals = sorted(existing_set - foundry_set)
    removal_rows = [existing_map[m] for m in removals if m in existing_map]

    result = {
        "excel": str(excel_path),
        "sheet": ws.title,
        "azure_subscription": account_info.get("id"),
        "azure_account_name": account_info.get("name"),
        "foundry_deployments_seen": len(foundry_set),
        "existing_upstream_in_sheet": len(existing_set),
        "to_append": len(additions),
        "append_items": [
            {"upstream_model_name": u, "ark_model_name": a, "provider_channel_type": p}
            for p, u, a in additions
        ],
        "to_remove": len(removals),
        "remove_items": removals,
        "prune_missing": args.prune_missing,
        "dry_run": args.dry_run,
    }

    if not args.dry_run:
        appended = append_model_rows(ws, h, additions) if additions else 0
        removed = prune_rows(ws, removal_rows) if args.prune_missing and removal_rows else 0
        if appended or removed:
            wb.save(excel_path)
            result["saved"] = True
        else:
            result["saved"] = False
        result["appended"] = appended
        result["removed"] = removed
    else:
        result["saved"] = False

    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
